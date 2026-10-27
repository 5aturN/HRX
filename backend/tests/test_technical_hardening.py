"""Регрессия прод-закалки: экранирование этикетки, даты отчётов, валидация заказов."""
import datetime as dt

import pytest
from sqlalchemy import select

from app.models.tables import AuditLog, Customer, Folder, Item, Order, StockMove

H = {"X-Telegram-Init-Data": "dev"}


async def _fid(session) -> int:
    f = (await session.execute(select(Folder).where(Folder.name == "Метизы"))).scalar_one()
    return f.id


async def _item(session, sku: str, kind: str = "standard", **kw) -> Item:
    it = Item(sku=sku, name=kw.pop("name", sku), kind=kind, folder_id=await _fid(session), **kw)
    session.add(it)
    await session.flush()
    return it


async def test_label_escapes_html(client, session):
    """Этикетка не должна исполнять HTML из артикула или названия."""
    it = await _item(session, '<img src=x onerror=alert(1)>', name='<script>alert("x")</script>')
    await session.commit()
    r = await client.get(f"/api/items/{it.id}/label", headers=H)
    assert r.status_code == 200, r.text
    body = r.text
    assert "<script>" not in body
    assert "<img src=x" not in body
    assert "&lt;script&gt;" in body
    assert "&lt;img src=x" in body


async def test_decrement_rejects_nonpositive_quantity(session):
    """Сервисное списание защищено и при вызове в обход Pydantic-схемы."""
    from app.services.moves import _decrement

    for qty in (0, -1):
        with pytest.raises(ValueError, match="положительным"):
            await _decrement(session, item_id=-1, qty=qty, message="тест")


async def test_moves_date_filters_apply_in_sql(client, session):
    """Фильтр дат отбирает только нужный день и включает его целиком."""
    it = await _item(session, "DATE-1", qty=10)
    session.add_all([
        StockMove(item_id=it.id, mtype="prihod", qty=1, reason="вчера",
                  created_at=dt.datetime(2026, 1, 1, 23, 59)),
        StockMove(item_id=it.id, mtype="prihod", qty=2, reason="граница",
                  created_at=dt.datetime(2026, 1, 2, 0, 0)),
        StockMove(item_id=it.id, mtype="brak", qty=3, reason="конец дня",
                  created_at=dt.datetime(2026, 1, 2, 23, 59)),
        StockMove(item_id=it.id, mtype="prihod", qty=4, reason="после",
                  created_at=dt.datetime(2026, 1, 3, 0, 0)),
    ])
    await session.commit()
    r = await client.get("/api/moves?date_from=2026-01-02&date_to=2026-01-02", headers=H)
    assert r.status_code == 200, r.text
    reasons = [m["reason"] for m in r.json()]
    assert reasons == ["конец дня", "граница"]

    r = await client.get("/api/moves?date_from=2026-01-02&date_to=2026-01-02&mtype=prihod", headers=H)
    assert [m["reason"] for m in r.json()] == ["граница"]


async def test_excel_date_filters_include_full_end_date(client, session):
    """Excel применяет границы даты в SQL, включая весь последний день."""
    from io import BytesIO

    from openpyxl import load_workbook

    it = await _item(session, "EXCEL-DATE-1", qty=10)
    session.add_all([
        StockMove(item_id=it.id, mtype="prihod", qty=1, reason="вчера",
                  created_at=dt.datetime(2026, 1, 1, 23, 59)),
        StockMove(item_id=it.id, mtype="prihod", qty=2, reason="начало",
                  created_at=dt.datetime(2026, 1, 2, 0, 0)),
        StockMove(item_id=it.id, mtype="brak", qty=3, reason="конец",
                  created_at=dt.datetime(2026, 1, 2, 23, 59)),
        StockMove(item_id=it.id, mtype="prihod", qty=4, reason="завтра",
                  created_at=dt.datetime(2026, 1, 3, 0, 0)),
    ])
    await session.commit()

    response = await client.get(
        "/api/report/excel?date_from=2026-01-02&date_to=2026-01-02", headers=H,
    )
    assert response.status_code == 200, response.text
    worksheet = load_workbook(BytesIO(response.content))["Движения"]
    assert [worksheet.cell(row, 6).value for row in range(2, worksheet.max_row + 1)] == ["конец", "начало"]


async def test_moves_bad_date_is_controlled(client):
    """Неверная дата — понятная ошибка клиента, а не 500."""
    r = await client.get("/api/moves?date_from=not-a-date", headers=H)
    assert r.status_code == 400, r.text


async def test_maximum_date_does_not_overflow(client):
    """Максимальная дата ISO не вызывает переполнение при построении границы."""
    r = await client.get("/api/moves?date_to=9999-12-31", headers=H)
    assert r.status_code == 200, r.text
    report = await client.get("/api/report/excel?date_to=9999-12-31", headers=H)
    assert report.status_code == 200, report.text


async def test_order_rejects_malformed_lines_without_partial_records(client, session):
    """Некорректная строка заказа не создаёт ни покупателя, ни заказ."""
    prod = await _item(session, "ORD-1", kind="product", qty=1)
    await session.commit()
    bad_payloads = [
        {"customer": "Новый клиент", "items": [{"product_id": prod.id, "need": "abc"}]},
        {"customer": "Новый клиент", "items": [{"product_id": prod.id, "need": None}]},
        {"customer": "Новый клиент", "items": [{"product_id": prod.id, "need": 0}]},
        {"customer": "Новый клиент", "items": [{"product_id": prod.id, "need": 1}, "not-a-dict"]},
        {"customer": "Новый клиент", "items": "not-a-list"},
        {"customer": "Новый клиент", "items": ""},
        {"customer": "Новый клиент", "items": 0},
        {"customer": "Новый клиент", "items": False},
        {"customer": "Новый клиент", "items": [{"product_id": True, "need": 1}]},
    ]
    for payload in bad_payloads:
        r = await client.post("/api/orders", json=payload, headers=H)
        assert r.status_code == 400, (payload, r.status_code, r.text)
    session.expire_all()
    assert (await session.execute(select(Customer).where(Customer.name == "Новый клиент"))).scalar_one_or_none() is None
    assert (await session.execute(select(Order))).scalars().all() == []


async def test_order_status_rejects_non_string(client, session):
    """Нестроковый статус — 400 без изменения заказа и без записи в аудит."""
    prod = await _item(session, "ORD-2", kind="product", qty=1)
    cust = Customer(name="Клиент")
    session.add(cust)
    await session.flush()
    order = Order(customer_id=cust.id, status="new")
    session.add(order)
    await session.commit()
    oid = order.id

    for value in (123, 1.5, [], {}, True):
        r = await client.post(f"/api/orders/{oid}/status", json={"status": value}, headers=H)
        assert r.status_code == 400, (value, r.status_code, r.text)

    session.expire_all()
    assert (await session.get(Order, oid)).status == "new"
    audit_rows = (await session.execute(select(AuditLog).where(AuditLog.action == "order_status"))).scalars().all()
    assert audit_rows == []


async def test_order_status_valid_still_works(client, session):
    """Корректный статус по-прежнему принимается и пишет аудит."""
    prod = await _item(session, "ORD-3", kind="product", qty=1)
    cust = Customer(name="Клиент 2")
    session.add(cust)
    await session.flush()
    order = Order(customer_id=cust.id, status="new")
    session.add(order)
    await session.commit()
    oid = order.id
    r = await client.post(f"/api/orders/{oid}/status", json={"status": "work"}, headers=H)
    assert r.status_code == 200, r.text
    session.expire_all()
    assert (await session.get(Order, oid)).status == "work"
    audit_rows = (await session.execute(select(AuditLog).where(
        AuditLog.action == "order_status", AuditLog.details.contains(f"Заказ №{oid}"),
    ))).scalars().all()
    assert len(audit_rows) == 1
