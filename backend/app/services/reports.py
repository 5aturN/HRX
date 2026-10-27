"""Сервис отчётов: Excel с двумя листами (Остатки + Движения), всё по-русски."""
import datetime as dt
import io
import logging
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from app.models.tables import Folder, Item, StockMove, User

log = logging.getLogger("hrx.reports")

MTYPE_RU = {
    "prihod": "Приход", "sborka": "Сборка", "brak": "Брак", "vozvrat": "Возврат",
    "vydacha_rd": "Выдача в проект", "sborka_podsborki": "Сборка подсборки",
    "oprihod_podsborki": "Оприходование",
}

async def build_excel(session: AsyncSession, date_from: dt.date | None, date_to: dt.date | None) -> bytes:
    """Собирает Excel: лист Остатки (на сейчас) + лист Движения (за период)."""
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Остатки"
    ws1.append(["Артикул", "Наименование", "Децимальный", "Папка", "Остаток", "Порог", "Статус"])
    for c in ws1[1]:
        c.font = Font(bold=True)
    items = (await session.execute(select(Item, Folder.name).join(Folder, Folder.id == Item.folder_id).order_by(Item.name))).all()
    for it, folder in items:
        low = it.notify_enabled and it.qty <= it.threshold
        ws1.append([it.sku, it.name, it.decimal or "", folder, it.qty, it.threshold, "Ниже порога" if low else "ОК"])
    for col in "ABCDEFG":
        ws1.column_dimensions[col].width = 22
    ws2 = wb.create_sheet("Движения")
    ws2.append(["Дата", "Тип", "Артикул", "Наименование", "Кол-во", "Причина", "Инициатор", "Заказ"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    stmt = select(StockMove, Item, User).join(Item, Item.id == StockMove.item_id).outerjoin(User, User.id == StockMove.user_id).order_by(StockMove.id.desc())
    if date_from:
        stmt = stmt.where(StockMove.created_at >= dt.datetime.combine(date_from, dt.time.min))
    if date_to and date_to < dt.date.max:
        stmt = stmt.where(StockMove.created_at < dt.datetime.combine(date_to + dt.timedelta(days=1), dt.time.min))
    rows = (await session.execute(stmt.limit(5000))).all()
    for m, it, u in rows:
        ws2.append([m.created_at.strftime("%d.%m.%Y %H:%M") if m.created_at else "",
                    MTYPE_RU.get(m.mtype, m.mtype), it.sku, it.name, m.qty,
                    m.reason or "", (u.first_name or u.username or "") if u else "",
                    m.order_id or "на склад"])
    for col in "ABCDEFGH":
        ws2.column_dimensions[col].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    log.info("Excel собран: остатков %s, движений %s", len(items), len(rows))
    return buf.getvalue()
